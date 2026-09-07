import io,base64
import pymupdf as fitz
from docx import Document
def decode_upload(content_base64:str,max_bytes:int)->bytes:
 data=base64.b64decode(content_base64,validate=True)
 if len(data)>max_bytes:raise ValueError("Resume exceeds configured size limit.")
 return data
def extract_text(filename:str,data:bytes)->str:
 if filename.lower().endswith(".pdf"):
  doc=fitz.open(stream=data,filetype="pdf");return "\n".join(p.get_text() for p in doc)
 if filename.lower().endswith(".docx"):
  doc=Document(io.BytesIO(data));chunks=[p.text for p in doc.paragraphs]
  for table in doc.tables:
   for row in table.rows:chunks.append(" | ".join(c.text for c in row.cells))
  return "\n".join(chunks)
 raise ValueError("Only PDF and DOCX resumes are supported.")
